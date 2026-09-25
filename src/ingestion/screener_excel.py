"""Screener.in Excel export parser.

How a Screener export is actually built (verified against a real export,
workbook version 2.1):

  Sheets: "Profit & Loss", "Quarters", "Balance Sheet", "Cash Flow",
          "Customization", "Data Sheet"

  The four presentation sheets contain ONLY FORMULAS that point at
  "Data Sheet", and Screener does not store cached results for them.
  openpyxl therefore returns None for every number on those sheets.
  All real values live in "Data Sheet", so that is what we parse.

  "Data Sheet" layout (column A = label, B..K = up to 10 periods):
    META           Face Value, Current Price, Market Capitalization
    PROFIT & LOSS  "Report Date" row of real Excel dates, then line items
    Quarters       "Report Date" row, then quarterly line items
    BALANCE SHEET  "Report Date" row, then balance sheet items
    CASH FLOW:     "Report Date" row, then CFO / CFI / CFF / Net
    PRICE:         Year-end share price, aligned to the P&L columns
    DERIVED:       "Adjusted Equity Shares in Cr" (bonus/split adjusted)

  Values are in ₹ crore, except share counts and prices.

  The export does NOT contain:
    - a Consolidated/Standalone marker (the caller must supply it)
    - a capital expenditure line (so FCF cannot be computed from it)
    - shareholding pattern data
"""

from __future__ import annotations

import calendar
import hashlib
import logging
import re
from dataclasses import dataclass, field
from datetime import date, datetime
from pathlib import Path
from typing import Any, Optional

import openpyxl

from src.ingestion.normalizer import (
    normalize_income_label,
    normalize_balance_label,
    normalize_cashflow_label,
    normalize_shareholding_label,
    clean_numeric,
)

logger = logging.getLogger(__name__)

STATEMENT_TYPES = ("Consolidated", "Standalone")
UNKNOWN_STATEMENT_TYPE = "Unknown"

# ---------------------------------------------------------------------------
# Data structures returned by the parser
# ---------------------------------------------------------------------------


@dataclass
class ParsedPeriod:
    """Represents one period column from the Screener export."""
    label: str          # e.g. "Mar 2026"
    fiscal_year: int    # e.g. 2026
    fiscal_quarter: Optional[int] = None   # only set for quarterly periods
    period_end_date: Optional[date] = None
    is_ttm: bool = False


@dataclass
class ParsedSheet:
    """One parsed statement with its periods and data rows."""
    sheet_name: str
    statement_type: str
    periods: list[ParsedPeriod]
    rows: dict[str, list[Optional[float]]]   # canonical_field → [val per period]
    unknown_labels: list[str] = field(default_factory=list)


@dataclass
class ParsedWorkbook:
    """Complete parsed output from one Screener Excel file."""
    file_path: str
    file_hash: str
    company_name: Optional[str]
    company_ticker: Optional[str]
    statement_type: str
    income_annual: Optional[ParsedSheet] = None
    balance_sheet: Optional[ParsedSheet] = None
    cash_flow: Optional[ParsedSheet] = None
    income_quarterly: Optional[ParsedSheet] = None
    shareholding: Optional[ParsedSheet] = None
    # Market data from the META block (current snapshot)
    current_price: Optional[float] = None
    market_cap: Optional[float] = None
    face_value: Optional[float] = None
    # Year-end share price, aligned to income_annual.periods
    year_end_prices: list[Optional[float]] = field(default_factory=list)
    source_layout: str = "unknown"   # "data_sheet" or "legacy_sheets"
    warnings: list[str] = field(default_factory=list)


# ---------------------------------------------------------------------------
# Period parsing
# ---------------------------------------------------------------------------

_MONTH_MAP = {
    "jan": 1, "feb": 2, "mar": 3, "apr": 4, "may": 5, "jun": 6,
    "jul": 7, "aug": 8, "sep": 9, "oct": 10, "nov": 11, "dec": 12,
}

# Indian FY: Apr–Jun = Q1, Jul–Sep = Q2, Oct–Dec = Q3, Jan–Mar = Q4
_QUARTER_END_MONTHS = {6: 1, 9: 2, 12: 3, 3: 4}


def _make_period(year: int, month: int, label: str, is_quarterly: bool) -> ParsedPeriod:
    """Build a ParsedPeriod for a period ending in (year, month)."""
    fiscal_year = year if month <= 3 else year + 1
    last_day = calendar.monthrange(year, month)[1]
    return ParsedPeriod(
        label=label,
        fiscal_year=fiscal_year,
        # Quarter is only meaningful for quarterly data. Setting it for
        # annual columns is what previously caused every annual year to be
        # skipped by the importer.
        fiscal_quarter=_QUARTER_END_MONTHS.get(month) if is_quarterly else None,
        period_end_date=date(year, month, last_day),
    )


def _parse_period_label(value: Any, is_quarterly: bool = False) -> Optional[ParsedPeriod]:
    """Parse a period header cell.

    Accepts real Excel dates (what Screener uses), and text such as
    'Mar 2026', 'Mar-26', 'Mar-2026', '2026-03-31', '2026' and 'TTM'.
    """
    if value is None:
        return None

    if isinstance(value, (datetime, date)):
        label = f"{calendar.month_abbr[value.month]} {value.year}"
        return _make_period(value.year, value.month, label, is_quarterly)

    s = str(value).strip()
    if not s:
        return None

    if "ttm" in s.lower():
        return ParsedPeriod(label=s, fiscal_year=0, is_ttm=True)

    # ISO date, optionally with a time part: "2026-03-31" / "2026-03-31 00:00:00"
    m = re.match(r"^(\d{4})-(\d{2})-(\d{2})", s)
    if m:
        year, month = int(m.group(1)), int(m.group(2))
        return _make_period(year, month, f"{calendar.month_abbr[month]} {year}", is_quarterly)

    # "Mar 2026", "Mar-2026", "Mar-26", "March 2026"
    m = re.match(r"^([A-Za-z]+)[\s\-']+(\d{2}|\d{4})$", s)
    if m:
        month = _MONTH_MAP.get(m.group(1).lower()[:3])
        if month is None:
            return None
        year = int(m.group(2))
        if year < 100:
            year += 2000
        return _make_period(year, month, f"{calendar.month_abbr[month]} {year}", is_quarterly)

    # Plain year "2026" → treat as March year-end
    m = re.match(r"^(\d{4})$", s)
    if m:
        return _make_period(int(m.group(1)), 3, f"Mar {m.group(1)}", is_quarterly)

    return None


# ---------------------------------------------------------------------------
# Data Sheet label maps (exact labels used by the Screener export)
# ---------------------------------------------------------------------------

_DS_PL = {
    "sales": "revenue",
    "raw material cost": "raw_material_cost",
    "change in inventory": "change_in_inventory",
    "power and fuel": "power_and_fuel",
    "other mfr. exp": "other_mfr_exp",
    "employee cost": "employee_cost",
    "selling and admin": "selling_and_admin",
    "other expenses": "other_expenses",
    "other income": "other_income",
    "depreciation": "depreciation",
    "interest": "interest",
    "profit before tax": "pbt",
    "tax": "tax",
    "net profit": "pat",
    "dividend amount": "dividend_amount",
}

_DS_QUARTERS = {
    "sales": "revenue",
    "expenses": "expenses",
    "other income": "other_income",
    "depreciation": "depreciation",
    "interest": "interest",
    "profit before tax": "pbt",
    "tax": "tax",
    "net profit": "pat",
    "operating profit": "operating_profit",
}

_DS_BALANCE = {
    "equity share capital": "equity_capital",
    "reserves": "reserves",
    "borrowings": "borrowings",
    "other liabilities": "other_liabilities",
    "net block": "fixed_assets",
    "capital work in progress": "cwip",
    "investments": "investments",
    "other assets": "other_assets",
    "receivables": "receivables",
    "inventory": "inventory",
    "cash & bank": "cash_and_equivalents",
    "no. of equity shares": "equity_shares_count",
    "new bonus shares": "new_bonus_shares",
    "face value": "face_value",
    # "Total" appears twice (liabilities side, then assets side) — handled
    # positionally in _parse_data_sheet.
}

_DS_CASHFLOW = {
    "cash from operating activity": "operating_cash_flow",
    "cash from investing activity": "investing_cash_flow",
    "cash from financing activity": "financing_cash_flow",
    "net cash flow": "net_cash_flow",
}

_DS_EXPENSE_COMPONENTS = (
    "raw_material_cost", "power_and_fuel", "other_mfr_exp",
    "employee_cost", "selling_and_admin", "other_expenses",
)

# Section markers in column A of the Data Sheet
_SECTION_MARKERS = {
    "meta": "meta",
    "profit & loss": "pl",
    "quarters": "quarters",
    "balance sheet": "balance",
    "cash flow": "cashflow",
    "price": "price",
    "derived": "derived",
}


def _norm(label: Any) -> str:
    s = str(label).strip().lower()
    s = re.sub(r"\s+", " ", s)
    return s.rstrip(":").strip()


# ---------------------------------------------------------------------------
# Main parser class
# ---------------------------------------------------------------------------

class ScreenerExcelParser:
    """Parse a Screener.in Excel export workbook."""

    def __init__(self) -> None:
        self._wb: Optional[openpyxl.Workbook] = None

    def parse(self, file_path: str | Path) -> ParsedWorkbook:
        """Parse the workbook and return structured data."""
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"File not found: {file_path}")

        file_hash = self._hash_file(path)
        logger.info("Parsing Screener export: %s (hash: %s)", path.name, file_hash[:8])

        self._wb = openpyxl.load_workbook(path, read_only=True, data_only=True)
        try:
            data_sheet = self._find_data_sheet()
            if data_sheet is not None:
                result = self._parse_data_sheet(data_sheet, path, file_hash)
            else:
                result = self._parse_legacy_sheets(path, file_hash)
        finally:
            self._wb.close()
        return result

    # ------------------------------------------------------------------
    # Data Sheet path (the real Screener format)
    # ------------------------------------------------------------------

    def _find_data_sheet(self):
        for name in self._wb.sheetnames:
            if name.strip().lower() == "data sheet":
                return self._wb[name]
        return None

    def _parse_data_sheet(self, ws, path: Path, file_hash: str) -> ParsedWorkbook:
        rows = [r for r in ws.iter_rows(values_only=True)]

        company_name: Optional[str] = None
        meta: dict[str, Optional[float]] = {}
        sections: dict[str, dict] = {}
        price_row: list[Any] = []
        adjusted_shares_row: list[Any] = []
        current: Optional[str] = None
        balance_totals_seen = 0

        for row in rows:
            if not row or row[0] is None:
                continue
            label = _norm(row[0])
            values = list(row[1:])

            if label == "company name":
                company_name = str(row[1]).strip() if len(row) > 1 and row[1] else None
                continue

            marker = next((v for k, v in _SECTION_MARKERS.items() if label == k), None)
            if marker is not None:
                current = marker
                if marker == "price":
                    price_row = values
                continue

            if current == "meta":
                if label == "current price":
                    meta["current_price"] = clean_numeric(row[1])
                elif label == "market capitalization":
                    meta["market_cap"] = clean_numeric(row[1])
                elif label == "face value":
                    meta["face_value"] = clean_numeric(row[1])
                continue

            if current == "derived":
                if label.startswith("adjusted equity shares"):
                    adjusted_shares_row = values
                continue

            if current not in ("pl", "quarters", "balance", "cashflow"):
                continue

            sec = sections.setdefault(current, {"periods": [], "cols": [], "rows": {}, "unknown": []})

            if label == "report date":
                is_q = current == "quarters"
                for idx, v in enumerate(values):
                    p = _parse_period_label(v, is_quarterly=is_q)
                    if p is not None:
                        sec["periods"].append(p)
                        sec["cols"].append(idx)
                continue

            label_map = {"pl": _DS_PL, "quarters": _DS_QUARTERS,
                         "balance": _DS_BALANCE, "cashflow": _DS_CASHFLOW}[current]
            if current == "balance" and label == "total":
                balance_totals_seen += 1
                canonical = "total_liabilities_and_equity" if balance_totals_seen == 1 else "total_assets"
            else:
                canonical = label_map.get(label)

            if canonical is None:
                sec["unknown"].append(str(row[0]).strip())
                continue
            sec["rows"][canonical] = [
                clean_numeric(values[c]) if c < len(values) else None for c in sec["cols"]
            ]

        result = ParsedWorkbook(
            file_path=str(path),
            file_hash=file_hash,
            company_name=company_name,
            company_ticker=None,
            statement_type=UNKNOWN_STATEMENT_TYPE,
            current_price=meta.get("current_price"),
            market_cap=meta.get("market_cap"),
            face_value=meta.get("face_value"),
            source_layout="data_sheet",
        )

        pl = sections.get("pl")
        if pl and pl["periods"]:
            self._derive_pl(pl["rows"], len(pl["periods"]))
            n = len(pl["periods"])
            adj = [clean_numeric(adjusted_shares_row[c]) if c < len(adjusted_shares_row) else None
                   for c in pl["cols"]]
            if any(v is not None for v in adj):
                pl["rows"]["weighted_avg_shares"] = adj
                pl["rows"]["basic_eps"] = [
                    round(p / s, 2) if p is not None and s else None
                    for p, s in zip(pl["rows"].get("pat", [None] * n), adj)
                ]
            result.year_end_prices = [
                clean_numeric(price_row[c]) if c < len(price_row) else None for c in pl["cols"]
            ]
            result.income_annual = self._to_sheet("Data Sheet / Profit & Loss", pl)

        q = sections.get("quarters")
        if q and q["periods"]:
            result.income_quarterly = self._to_sheet("Data Sheet / Quarters", q)

        bs = sections.get("balance")
        if bs and bs["periods"]:
            self._derive_balance(bs["rows"], len(bs["periods"]))
            result.balance_sheet = self._to_sheet("Data Sheet / Balance Sheet", bs)

        cf = sections.get("cashflow")
        if cf and cf["periods"]:
            result.cash_flow = self._to_sheet("Data Sheet / Cash Flow", cf)
            result.warnings.append(
                "Screener exports contain no capital expenditure line, so Free Cash Flow "
                "is not available from this import."
            )

        for sheet in (result.income_annual, result.balance_sheet,
                      result.cash_flow, result.income_quarterly):
            if sheet and sheet.unknown_labels:
                result.warnings.append(
                    f"[{sheet.sheet_name}] Unrecognized labels: "
                    + ", ".join(sheet.unknown_labels[:10])
                )
        return result

    @staticmethod
    def _to_sheet(name: str, sec: dict) -> ParsedSheet:
        return ParsedSheet(
            sheet_name=name,
            statement_type=UNKNOWN_STATEMENT_TYPE,
            periods=sec["periods"],
            rows=sec["rows"],
            unknown_labels=sec["unknown"],
        )

    @staticmethod
    def _derive_pl(rows: dict, n: int) -> None:
        """Replicate Screener's own P&L formulas.

        Expenses         = RM + Power + Other Mfr + Employee + Selling + Other − Change in Inventory
        Operating Profit = Sales − Expenses   (this is Screener's EBITDA, excl. other income)
        EBIT             = PBT + Interest     (Screener's ROCE convention; includes other income)
        """
        def col(k: str) -> list[Optional[float]]:
            return rows.get(k, [None] * n)

        expenses: list[Optional[float]] = []
        for i in range(n):
            parts = [col(k)[i] for k in _DS_EXPENSE_COMPONENTS]
            if all(p is None for p in parts):
                expenses.append(None)
                continue
            total = sum(p for p in parts if p is not None)
            total -= col("change_in_inventory")[i] or 0.0
            expenses.append(total)
        rows["expenses"] = expenses

        rev = col("revenue")
        rows["operating_profit"] = [
            r - e if r is not None and e is not None else None for r, e in zip(rev, expenses)
        ]
        rows["ebitda"] = list(rows["operating_profit"])
        rows["ebit"] = [
            p + i if p is not None and i is not None else None
            for p, i in zip(col("pbt"), col("interest"))
        ]

    @staticmethod
    def _derive_balance(rows: dict, n: int) -> None:
        def col(k: str) -> list[Optional[float]]:
            return rows.get(k, [None] * n)

        rows["shareholders_equity"] = [
            c + r if c is not None and r is not None else None
            for c, r in zip(col("equity_capital"), col("reserves"))
        ]
        rows["total_liabilities"] = [
            (b or 0.0) + (o or 0.0) if (b is not None or o is not None) else None
            for b, o in zip(col("borrowings"), col("other_liabilities"))
        ]

    # ------------------------------------------------------------------
    # Legacy path: one sheet per statement (only used without a Data Sheet)
    # ------------------------------------------------------------------

    def _parse_legacy_sheets(self, path: Path, file_hash: str) -> ParsedWorkbook:
        result = ParsedWorkbook(
            file_path=str(path),
            file_hash=file_hash,
            company_name=self._detect_company(),
            company_ticker=None,
            statement_type=UNKNOWN_STATEMENT_TYPE,
            source_layout="legacy_sheets",
        )
        result.income_annual = self._parse_sheet(
            ["Profit & Loss", "P&L", "Profit and Loss"], normalize_income_label, is_quarterly=False)
        result.balance_sheet = self._parse_sheet(
            ["Balance Sheet", "BalanceSheet"], normalize_balance_label, is_quarterly=False)
        result.cash_flow = self._parse_sheet(
            ["Cash Flow", "Cashflow"], normalize_cashflow_label, is_quarterly=False)
        result.income_quarterly = self._parse_sheet(
            ["Quarters", "Quarterly", "Quarterly Results"], normalize_income_label, is_quarterly=True)
        result.shareholding = self._parse_sheet(
            ["Shareholding", "Shareholding Pattern"], normalize_shareholding_label, is_quarterly=True)

        for sheet in (result.income_annual, result.balance_sheet, result.cash_flow,
                      result.income_quarterly, result.shareholding):
            if sheet is None:
                continue
            if sheet.rows and all(v is None for vals in sheet.rows.values() for v in vals):
                result.warnings.append(
                    f"[{sheet.sheet_name}] Every value is empty. This sheet probably contains "
                    "formulas without saved results. Re-export from Screener without editing, "
                    "or open and save the file in Excel first."
                )
            if sheet.unknown_labels:
                result.warnings.append(
                    f"[{sheet.sheet_name}] Unrecognized labels: " + ", ".join(sheet.unknown_labels[:10])
                )
        return result

    def _detect_company(self) -> Optional[str]:
        for sheet_name in self._wb.sheetnames:
            val = self._wb[sheet_name].cell(1, 1).value
            if isinstance(val, str) and len(val) > 1 and not val.startswith("="):
                return re.sub(r"\s*\(.*?\)", "", val).strip()
        return None

    def _find_sheet(self, candidates: list[str]):
        sheet_map = {s.strip().lower(): s for s in self._wb.sheetnames}
        for candidate in candidates:
            actual = sheet_map.get(candidate.lower())
            if actual:
                return self._wb[actual]
        return None

    def _parse_sheet(self, candidates: list[str], normalize_fn: Any,
                     is_quarterly: bool) -> Optional[ParsedSheet]:
        ws = self._find_sheet(candidates)
        if ws is None:
            return None
        all_rows = list(ws.iter_rows(values_only=True))
        header_idx = None
        for i, row in enumerate(all_rows[:10]):
            if sum(1 for c in row[1:] if _parse_period_label(c, is_quarterly)) >= 2:
                header_idx = i
                break
        if header_idx is None:
            logger.warning("No period header found in sheet %s", ws.title)
            return None

        periods, cols = [], []
        for idx, v in enumerate(all_rows[header_idx]):
            if idx == 0:
                continue
            p = _parse_period_label(v, is_quarterly)
            if p is not None:
                periods.append(p)
                cols.append(idx)

        rows: dict[str, list[Optional[float]]] = {}
        unknown: list[str] = []
        for row in all_rows[header_idx + 1:]:
            if not row or row[0] is None or not str(row[0]).strip():
                continue
            canonical = normalize_fn(str(row[0]))
            if canonical is None:
                unknown.append(str(row[0]).strip())
                continue
            if canonical not in rows:
                rows[canonical] = [clean_numeric(row[c]) if c < len(row) else None for c in cols]
        return ParsedSheet(ws.title, UNKNOWN_STATEMENT_TYPE, periods, rows, unknown)

    @staticmethod
    def _hash_file(path: Path) -> str:
        return compute_file_hash(path)


def compute_file_hash(file_path: str | Path) -> str:
    """SHA-256 of a file."""
    h = hashlib.sha256()
    with open(file_path, "rb") as f:
        for chunk in iter(lambda: f.read(65536), b""):
            h.update(chunk)
    return h.hexdigest()
